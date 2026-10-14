# -*- coding: utf-8 -*-
"""
=============================================================================
GaiaOps ESG - FIAP | Navegando pelo mundo DevOps
tools/gerar_documento.py

Gera a documentacao tecnica da entrega:  docs/GaiaOps-ESG-DevOps.docx
(converta para PDF com tools/gerar_pdf.ps1).

Caracteristicas:
  1. Os trechos de codigo (Dockerfile, docker-compose.yml, workflow) sao LIDOS
     DOS ARQUIVOS REAIS do projeto - o documento nunca fica defasado.
  2. Os prints sao inseridos automaticamente: para cada print previsto, o
     gerador procura docs/prints/<nome>.png. Se existir, entra a imagem; se
     nao, entra um quadro laranja dizendo o que capturar.
  3. O checklist final marca como concluidos os itens comprovados pelos
     arquivos presentes (ex.: prints de staging e producao).

Antes de gerar a versao final, preencha INTEGRANTES e REPOSITORIO abaixo.

Uso:  py tools/gerar_documento.py
=============================================================================
"""

import datetime
import os
import sys

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:
    print("FALTA A BIBLIOTECA python-docx. Instale com:  py -m pip install python-docx")
    sys.exit(1)

# =============================================================================
# PREENCHA ANTES DA ENTREGA
# =============================================================================
INTEGRANTES = [
    ("Nome do integrante", "RM00000"),
]
REPOSITORIO = "https://github.com/JotaDiasDevs/gaiaops-esg"

# =============================================================================
RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DIR_PRINTS = os.path.join(RAIZ, "docs", "prints")
SAIDA_DOCX = os.path.join(RAIZ, "docs", "GaiaOps-ESG-DevOps.docx")

# Paleta (a mesma da documentacao da Fase 5)
VERDE_ESCURO = RGBColor(0x1B, 0x4D, 0x3E)
VERDE_MEDIO = RGBColor(0x2D, 0x6A, 0x4F)
CINZA_TEXTO = RGBColor(0x33, 0x33, 0x33)
CINZA_CODIGO = "F5F5F5"
VERDE_CLARO = "E8F3EE"
LARANJA_FUNDO = "FFF4E5"
LARANJA_BORDA = "E67E22"

# Secoes de primeiro nivel - usadas no sumario e nos titulos.
SECOES = [
    "1. O projeto",
    "2. Pipeline de CI/CD",
    "3. Docker: arquitetura, imagem e comandos",
    "4. Testes automatizados",
    "5. Evidências: pipeline em execução",
    "6. Evidências: ambientes de staging e produção",
    "7. Desafios encontrados e como foram resolvidos",
    "8. Tecnologias utilizadas",
    "9. Checklist de entrega",
]


# =============================================================================
# LEITURA DOS ARQUIVOS DO PROJETO
# =============================================================================

def ler_arquivo(relativo):
    with open(os.path.join(RAIZ, relativo), encoding="utf-8") as f:
        return f.read()


def sem_comentarios(texto, marcador="#"):
    """Remove linhas so de comentario e linhas em branco repetidas."""
    linhas = []
    for linha in texto.splitlines():
        if linha.strip().startswith(marcador) and not linha.strip().startswith("#!"):
            continue
        if not linha.strip() and (not linhas or not linhas[-1].strip()):
            continue
        linhas.append(linha)
    return "\n".join(linhas).strip()


def trecho(texto, inicio, fim=None):
    """Recorta de 'inicio' (inclusive) ate 'fim' (exclusive)."""
    i = texto.index(inicio)
    j = texto.index(fim, i + len(inicio)) if fim else len(texto)
    return texto[i:j].rstrip()


def existe_print(nome):
    return any(os.path.exists(os.path.join(DIR_PRINTS, nome + ext)) for ext in (".png", ".jpg", ".jpeg"))


def caminho_print(nome):
    for ext in (".png", ".jpg", ".jpeg"):
        caminho = os.path.join(DIR_PRINTS, nome + ext)
        if os.path.exists(caminho):
            return caminho
    return None


# =============================================================================
# HELPERS DE FORMATACAO
# =============================================================================

def _sombrear(elemento_pr, cor_hex):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor_hex)
    elemento_pr.append(shd)


def _borda_paragrafo(paragrafo, cor, tamanho=6):
    pPr = paragrafo._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    for lado in ("top", "bottom", "left", "right"):
        el = OxmlElement("w:" + lado)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(tamanho))
        el.set(qn("w:space"), "6")
        el.set(qn("w:color"), cor)
        pBdr.append(el)
    pPr.append(pBdr)


class Documento(object):

    def __init__(self):
        self.doc = Document()
        secao = self.doc.sections[0]
        secao.page_height, secao.page_width = Cm(29.7), Cm(21.0)
        secao.left_margin = secao.right_margin = Cm(2.2)
        secao.top_margin = secao.bottom_margin = Cm(2.0)

        normal = self.doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        normal.font.color.rgb = CINZA_TEXTO
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15

        for nivel, tamanho, cor in ((1, 18, VERDE_ESCURO), (2, 13.5, VERDE_MEDIO), (3, 12, VERDE_MEDIO)):
            estilo = self.doc.styles["Heading %d" % nivel]
            estilo.font.name = "Calibri"
            estilo.font.size = Pt(tamanho)
            estilo.font.bold = True
            estilo.font.color.rgb = cor
            estilo.paragraph_format.space_before = Pt(14 if nivel == 1 else 10)
            estilo.paragraph_format.space_after = Pt(6)

        self._rodape()

    def _rodape(self):
        p = self.doc.sections[0].footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("GaiaOps ESG · Cidades ESG Inteligentes · FIAP - Navegando pelo mundo DevOps · página ")
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        for tipo, texto in (("begin", None), (None, "PAGE"), ("end", None)):
            run = p.add_run()
            run.font.size = Pt(8)
            if tipo:
                fld = OxmlElement("w:fldChar")
                fld.set(qn("w:fldCharType"), tipo)
                run._r.append(fld)
            else:
                instr = OxmlElement("w:instrText")
                instr.set(qn("xml:space"), "preserve")
                instr.text = texto
                run._r.append(instr)

    # --- blocos -----------------------------------------------------------------

    def quebra_pagina(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def h1(self, texto, nova_pagina=True):
        if nova_pagina:
            self.quebra_pagina()
        self.doc.add_heading(texto, level=1)

    def h2(self, texto):
        self.doc.add_heading(texto, level=2)

    def h3(self, texto):
        self.doc.add_heading(texto, level=3)

    def _texto_rico(self, paragrafo, texto):
        """Suporta **negrito**, *italico* e `codigo` inline (sem aninhamento)."""
        partes = []
        buffer, modo, i = "", None, 0
        while i < len(texto):
            if texto.startswith("**", i) and modo in (None, "b"):
                partes.append((buffer, modo))
                buffer, modo = "", (None if modo == "b" else "b")
                i += 2
                continue
            if texto[i] == "*" and modo in (None, "i"):
                partes.append((buffer, modo))
                buffer, modo = "", (None if modo == "i" else "i")
                i += 1
                continue
            if texto[i] == "`" and modo in (None, "c"):
                partes.append((buffer, modo))
                buffer, modo = "", (None if modo == "c" else "c")
                i += 1
                continue
            buffer += texto[i]
            i += 1
        partes.append((buffer, modo))
        for conteudo, estilo in partes:
            if not conteudo:
                continue
            run = paragrafo.add_run(conteudo)
            if estilo == "b":
                run.bold = True
            elif estilo == "i":
                run.italic = True
            elif estilo == "c":
                run.font.name = "Consolas"
                run.font.size = Pt(9.5)
                run.font.color.rgb = VERDE_ESCURO

    def par(self, texto):
        p = self.doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        self._texto_rico(p, texto)
        return p

    def bullets(self, itens):
        for item in itens:
            p = self.doc.add_paragraph(style="List Bullet")
            self._texto_rico(p, item)

    def numerados(self, itens):
        for item in itens:
            p = self.doc.add_paragraph(style="List Number")
            self._texto_rico(p, item)

    def codigo(self, texto, legenda=None):
        if legenda:
            p = self.doc.add_paragraph()
            r = p.add_run(legenda)
            r.italic = True
            r.font.size = Pt(9)
            p.paragraph_format.space_after = Pt(2)
        tabela = self.doc.add_table(rows=1, cols=1)
        tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
        celula = tabela.rows[0].cells[0]
        _sombrear(celula._tc.get_or_add_tcPr(), CINZA_CODIGO)
        celula.paragraphs[0].text = ""
        for i, linha in enumerate(texto.splitlines()):
            p = celula.paragraphs[0] if i == 0 else celula.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(linha if linha else " ")
            r.font.name = "Consolas"
            r.font.size = Pt(8.5)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def nota(self, texto):
        p = self.doc.add_paragraph()
        _sombrear(p._p.get_or_add_pPr(), VERDE_CLARO)
        _borda_paragrafo(p, "2D6A4F", 4)
        self._texto_rico(p, texto)

    def tabela(self, cabecalho, linhas, larguras=None):
        t = self.doc.add_table(rows=1, cols=len(cabecalho))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, titulo in enumerate(cabecalho):
            celula = t.rows[0].cells[i]
            _sombrear(celula._tc.get_or_add_tcPr(), "1B4D3E")
            celula.paragraphs[0].text = ""
            r = celula.paragraphs[0].add_run(titulo)
            r.bold = True
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        for linha in linhas:
            celulas = t.add_row().cells
            for i, valor in enumerate(linha):
                celulas[i].paragraphs[0].text = ""
                self._texto_rico(celulas[i].paragraphs[0], str(valor))
                for r in celulas[i].paragraphs[0].runs:
                    r.font.size = Pt(9.5)
        if larguras:
            # O LibreOffice le a largura do <w:tblGrid> (colunas), o Word a da
            # celula: as duas sao definidas, e o ajuste automatico e desligado.
            t.autofit = False
            for i, largura in enumerate(larguras):
                t.columns[i].width = Cm(largura)
                for celula in t.columns[i].cells:
                    celula.width = Cm(largura)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def print_tela(self, nome, legenda, o_que_capturar):
        """Insere docs/prints/<nome>.png ou um quadro indicando o print pendente."""
        caminho = caminho_print(nome)
        if caminho:
            self.doc.add_picture(caminho, width=Cm(16.5))
            self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            p = self.doc.add_paragraph()
            _sombrear(p._p.get_or_add_pPr(), LARANJA_FUNDO)
            _borda_paragrafo(p, LARANJA_BORDA, 12)
            r = p.add_run("PRINT PENDENTE: docs/prints/%s.png\n" % nome)
            r.bold = True
            r.font.color.rgb = RGBColor(0xB3, 0x5C, 0x00)
            p.add_run(o_que_capturar).font.size = Pt(10)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(legenda)
        r.italic = True
        r.font.size = Pt(9)

    def salvar(self):
        self.doc.save(SAIDA_DOCX)


# =============================================================================
# CONTEUDO
# =============================================================================

def capa(d):
    for _ in range(5):
        d.doc.add_paragraph()
    p = d.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("GaiaOps ESG")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(34), VERDE_ESCURO

    p = d.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Cidades ESG Inteligentes")
    r.font.size, r.font.color.rgb = Pt(18), VERDE_MEDIO

    p = d.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Pipeline de CI/CD, containerização e orquestração\nde uma API Java Spring Boot + MongoDB")
    r.font.size = Pt(13)

    for _ in range(3):
        d.doc.add_paragraph()
    p = d.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("FIAP · Navegando pelo mundo DevOps")
    r.bold, r.font.size = True, Pt(12)

    d.doc.add_paragraph()
    d.tabela(["Integrante", "RM"], INTEGRANTES, larguras=[11, 4])

    p = d.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Repositório: ").bold = True
    p.add_run(REPOSITORIO)
    p = d.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(datetime.date.today().strftime("%d/%m/%Y"))


def sumario(d):
    d.quebra_pagina()
    d.doc.add_heading("Sumário", level=1)
    for secao in SECOES:
        p = d.doc.add_paragraph(secao)
        p.paragraph_format.space_after = Pt(4)


def secao_projeto(d):
    d.h1(SECOES[0])
    d.h2("1.1 Ponto de partida")
    d.par("O GaiaOps é a plataforma de gestão ESG do município fictício de **Vila Nova do Horizonte** "
          "(cerca de 90 mil habitantes, quatro distritos). Ele surgiu na Fase 5 como uma base MongoDB com "
          "cinco collections: `equipamentos`, `medicoes`, `residuos`, `programas_sociais` e `conformidade`. "
          "Elas cobrem os três pilares ESG: telemetria ambiental, coleta de resíduos, programas sociais e "
          "governança municipal.")
    d.par("Para esta fase, a base foi transformada em uma **API Java 21 + Spring Boot 3.5** e adaptada para "
          "um ambiente de produção simulado. A API reaproveita integralmente o trabalho anterior: os mesmos "
          "validadores `$jsonSchema` e índices são aplicados automaticamente no container do banco, e os "
          "100 documentos da Fase 5 entram como carga inicial.")

    d.h2("1.2 A aplicação")
    d.tabela(["Recurso", "Descrição"], [
        ["`/api/equipamentos`", "CRUD completo. Exclusão lógica e atualização com `$set`, que preserva os atributos específicos de cada tipo"],
        ["`/api/medicoes`", "Ingestão de leituras de sensores com **alertas automáticos** (IQAr, ruído, perdas de água, nível de córregos)"],
        ["`/api/indicadores`", "Painel ESG dos três pilares, calculado por agregações no MongoDB"],
        ["`/api/residuos`, `/api/programas-sociais`, `/api/conformidade`", "Consultas, resumo de resíduos e licenças vencidas ou a vencer"],
        ["`/actuator/health`, `/actuator/info`", "Saúde (aplicação + banco) e versão/ambiente/commit implantado"],
        ["`/` e `/swagger-ui.html`", "Página inicial com o painel ao vivo e documentação interativa"],
    ], larguras=[5.5, 11])

    d.h2("1.3 Visão geral da solução")
    d.codigo(
        "  desenvolvedor --git push--> GitHub --dispara--> GitHub Actions\n"
        "                                                     |\n"
        "     1. Build --> 2. Testes --> 3. Imagem Docker ----+--push--> GHCR (registry)\n"
        "                                                     |              |\n"
        "                          4. Deploy staging <--------+-- pull ------+\n"
        "                               | smoke test OK                      |\n"
        "                          [aprovação manual]                        |\n"
        "                               v                                    |\n"
        "                          5. Deploy produção <------- pull ---------+\n"
        "\n"
        "  Cada ambiente (Docker Compose):   [api :8080] --rede backend--> [mongo 8.0] --- (volume)",
        "Figura 1 - Fluxo do código até produção")


def secao_pipeline(d):
    d.h1(SECOES[1])
    d.h2("2.1 Ferramenta: GitHub Actions")
    d.par("O pipeline foi implementado no **GitHub Actions** (`.github/workflows/ci-cd.yml`). A escolha "
          "considerou a integração nativa com o repositório, os **Environments** com aprovação manual e "
          "segredos por ambiente, o **GitHub Container Registry (GHCR)** para publicar a imagem e os runners "
          "Linux com Docker, que permitem rodar testes com containers e fazer o deploy com Docker Compose sem "
          "infraestrutura adicional.")

    d.h2("2.2 Etapas")
    d.tabela(["#", "Job", "O que faz", "Quando"], [
        ["1", "Build", "Compila com Maven (Java 21 Temurin, cache de dependências), empacota o `.jar` e o publica como artefato", "push e PR"],
        ["2", "Testes automatizados", "`mvn verify`: testes unitários (surefire) e de integração com MongoDB real via Testcontainers (failsafe). Publica relatórios, cobertura JaCoCo e um resumo na página da execução", "push e PR"],
        ["3", "Imagem Docker", "Build da imagem pelo `Dockerfile` com cache de camadas e push no GHCR com a tag do commit e `latest`", "push na main"],
        ["4", "Deploy staging", "`deploy.sh staging` (pull da imagem + `docker compose up --wait`), smoke test, print da página, coleta de logs e tag `staging` na imagem", "após o 3"],
        ["5", "Deploy produção", "Aguarda aprovação manual. Repete o deploy com a mesma imagem e a configuração de produção, roda o smoke test e aplica a tag `production`", "após aprovação"],
    ], larguras=[0.7, 3.2, 10.3, 2.3])

    d.h2("2.3 Lógica do pipeline")
    d.bullets([
        "**Portões em sequência.** Cada job declara `needs` do anterior. Um teste quebrado não gera imagem, "
        "e uma imagem que falha no smoke test de staging nunca chega à produção.",
        "**Build once, deploy many.** A imagem é construída uma única vez (job 3) e identificada pelo hash do "
        "commit. Staging e produção baixam exatamente a mesma imagem. Entre os ambientes muda apenas a "
        "configuração (`deploy/staging.env` e `deploy/production.env`).",
        "**Pull requests** executam apenas build e testes, como validação antes do merge.",
        "**Aprovação manual.** O job de produção usa o environment `production`, configurado com "
        "*Required reviewers*. A execução para e espera um revisor aprovar.",
        "**Segredos fora do código.** As senhas do MongoDB vêm dos GitHub Secrets de cada environment. Sem "
        "eles, o `deploy.sh` gera senhas aleatórias no momento do deploy. Em ambos os casos elas são "
        "mascaradas no log (`::add-mask::`).",
        "**Promoção rastreável.** Após cada deploy bem-sucedido, a imagem recebe a tag do ambiente (`staging` e "
        "`production`) no GHCR. A qualquer momento é possível saber qual commit está em cada ambiente, "
        "também visível em `/actuator/info`.",
        "**Mesmos scripts no pipeline e na máquina do desenvolvedor.** O workflow só chama `deploy/deploy.sh` "
        "e `deploy/smoke-test.sh`, que também funcionam localmente.",
    ])

    d.h2("2.4 Smoke test pós-deploy")
    d.par("Após subir cada ambiente, `deploy/smoke-test.sh` faz requisições reais e só deixa o pipeline seguir "
          "se todas as verificações passarem:")
    d.tabela(["#", "Verificação"], [
        ["1", "`/actuator/health/readiness` = UP (aplicação pronta e MongoDB conectado)"],
        ["2", "`/actuator/info` informa o ambiente correto (`staging` ou `producao`)"],
        ["3", "`/actuator/info` informa o commit que está sendo implantado"],
        ["4", "`/api/indicadores` devolve o painel ESG calculado"],
        ["5", "`/api/equipamentos` devolve os dados da carga inicial"],
        ["6", "Cadastro inválido é recusado com HTTP 400 (validação ativa)"],
        ["7", "A página inicial é publicada"],
        ["8", "Configuração do ambiente aplicada: Swagger disponível em staging e **desabilitado em produção**"],
    ], larguras=[0.8, 15.7])

    d.h2("2.5 Diferenças entre os ambientes")
    d.tabela(["", "Staging", "Produção"], [
        ["Perfil Spring", "`staging`", "`prod`"],
        ["Porta publicada", "8081", "8082"],
        ["Swagger UI", "habilitado", "desabilitado"],
        ["Detalhes no health check", "sim (Mongo, readiness)", "não (só UP/DOWN)"],
        ["Nível de log da aplicação", "DEBUG", "INFO"],
        ["Limite de memória da API", "768 MB", "1 GB"],
        ["Aprovação para deploy", "automática", "manual (required reviewers)"],
    ], larguras=[5.5, 5.5, 5.5])

    d.h2("2.6 Onde os ambientes rodam")
    d.par("Os jobs de deploy rodam por padrão nos **runners hospedados do GitHub**: máquinas Linux com Docker "
          "criadas para cada job e descartadas ao final. O deploy é completo e verificável. A imagem é baixada "
          "do registry, os containers sobem com rede, volume e variáveis de ambiente, o smoke test valida o "
          "ambiente e as evidências (respostas da API, print da página, estado e logs dos containers) ficam "
          "salvas como artefatos da execução. O ambiente, porém, existe apenas durante o job.")
    d.nota("Para ambientes persistentes, o workflow aceita um **self-hosted runner**: basta registrar uma "
           "máquina Linux com Docker e criar a variável `DEPLOY_RUNNER`. Os mesmos jobs passam a implantar "
           "nessa máquina, onde staging (8081) e produção (8082) permanecem no ar, sem alterar nenhum script.")

    d.h2("2.7 Trecho do workflow: job de produção")
    workflow = ler_arquivo(".github/workflows/ci-cd.yml")
    d.codigo(trecho(workflow, "  deploy-production:", "      - name: Captura de tela"),
             "Listagem 1 - .github/workflows/ci-cd.yml (job 5)")


def secao_docker(d):
    d.h1(SECOES[2])
    d.h2("3.1 Arquitetura dos containers")
    d.codigo(
        "  host (runner do GitHub ou máquina local)\n"
        "  |\n"
        "  |  porta 8081 (staging) / 8082 (produção) / 8080 (local)\n"
        "  v\n"
        "  +----------------------- projeto Compose gaiaops-<ambiente> -------------------+\n"
        "  |                                                                               |\n"
        "  |   rede frontend          rede backend (internal: sem acesso externo)          |\n"
        "  |  ---------------+  +-------------------------------------------+            |\n"
        "  |                 |  |                                           |            |\n"
        "  |        +--------v--v--------+                 +----------------v--------+   |\n"
        "  |        |  api               |   MONGODB_URI   |  mongo (mongo:8.0)      |   |\n"
        "  |        |  Spring Boot :8080 | --------------> |  usuário gaiaops_app    |   |\n"
        "  |        |  usuário 10001     |   (readWrite)   |  sem porta publicada    |   |\n"
        "  |        +--------------------+                 +-----------+-------------+   |\n"
        "  |          healthcheck: readiness                          |                 |\n"
        "  |                                           volume mongo-data   init: 01-setup.js\n"
        "  +-------------------------------------------------------------------------------+",
        "Figura 2 - Containers, redes e volume de um ambiente")
    d.bullets([
        "**Dois serviços:** `mongo` e `api`. A API só inicia quando o MongoDB está *healthy* "
        "(`depends_on: condition: service_healthy`).",
        "**Duas redes:** `frontend`, onde a API publica a porta, e `backend`, marcada como `internal`, que liga "
        "só a API ao banco. O MongoDB **não expõe porta** para fora do Docker.",
        "**Volume nomeado** `mongo-data`: os dados sobrevivem a reinícios e a novos deploys.",
        "**Script de inicialização** `mongo/init/01-setup.js`: roda uma única vez, na criação do volume, e cria "
        "as collections com os validadores, os índices da Fase 5 e um usuário da aplicação com permissão só no "
        "banco dela (menor privilégio).",
        "**Variáveis de ambiente:** tudo que muda entre ambientes vem de `.env` (local) ou "
        "`deploy/<ambiente>.env`. Senhas obrigatórias usam `${VAR:?}`, então o Compose se recusa a subir sem elas.",
    ])

    d.h2("3.2 Dockerfile")
    d.codigo(sem_comentarios(ler_arquivo("Dockerfile")), "Listagem 2 - Dockerfile (sem os comentários)")
    d.tabela(["Estratégia", "Benefício"], [
        ["Multi-stage build", "Imagem final só com JRE + `.jar`. Maven, JDK e código-fonte ficam no estágio de build"],
        ["`pom.xml` antes do código", "Dependências em camada própria, reaproveitada do cache quando só o código muda"],
        ["Base `eclipse-temurin:21-jre-alpine`", "Runtime mínimo e mantido, com menor superfície de ataque"],
        ["Usuário sem privilégios (UID 10001)", "O processo não roda como root no container"],
        ["`JAVA_TOOL_OPTIONS`", "Heap proporcional ao limite de memória do container; encerra em OutOfMemoryError"],
        ["`HEALTHCHECK` no readiness", "*healthy* só com a API pronta e o MongoDB respondendo, base do `compose up --wait`"],
        ["`ENTRYPOINT` em formato exec", "A JVM é o PID 1 e recebe o SIGTERM, garantindo o encerramento gracioso"],
        ["`.dockerignore`", "Só `pom.xml` e `src/` vão para o build; `.env`, `.git` e `target` ficam fora"],
    ], larguras=[5.5, 11])

    d.h2("3.3 docker-compose.yml")
    d.codigo(sem_comentarios(ler_arquivo("docker-compose.yml")), "Listagem 3 - docker-compose.yml (sem os comentários)")

    d.h2("3.4 Imagem criada")
    d.tabela(["Item", "Valor"], [
        ["Registry", "GitHub Container Registry (GHCR)"],
        ["Nome", "`ghcr.io/<usuario>/gaiaops-esg`"],
        ["Tags", "hash completo do commit (imutável), `latest`, `staging` e `production`"],
        ["Base", "`eclipse-temurin:21-jre-alpine`"],
        ["Conteúdo", "JRE 21 + `app.jar` (Spring Boot), usuário 10001, porta 8080, healthcheck"],
        ["Metadados", "labels OCI (origem, revisão, data de criação) gerados pelo `docker/metadata-action`"],
    ], larguras=[3.5, 13])

    d.h2("3.5 Comandos utilizados")
    d.tabela(["Comando", "Finalidade"], [
        ["`docker compose up -d --build`", "Constrói a imagem e sobe API + MongoDB localmente"],
        ["`docker compose ps`", "Estado e healthcheck dos serviços"],
        ["`docker compose logs -f api`", "Acompanha o log da API"],
        ["`docker compose down` / `down -v`", "Para o ambiente (mantendo ou apagando o volume)"],
        ["`./deploy/deploy.sh staging|production`", "Sobe um ambiente: `docker compose -p gaiaops-<amb> --env-file deploy/<amb>.env up -d --wait`"],
        ["`./deploy/smoke-test.sh staging|production`", "Valida o ambiente implantado (8 verificações)"],
        ["`docker compose pull api`", "No pipeline: baixa do GHCR a imagem já testada"],
        ["`docker buildx imagetools create --tag ...:staging`", "Promove a imagem (aplica a tag do ambiente)"],
        ["`./mvnw verify`", "Build + testes unitários e de integração"],
    ], larguras=[7, 9.5])


def secao_testes(d):
    d.h1(SECOES[3])
    d.par("Os testes rodam no job 2, com `./mvnw verify`, e qualquer falha interrompe o pipeline.")
    d.tabela(["Classe", "Tipo", "O que verifica", "Testes"], [
        ["`RegrasDeAlertaTest`", "Unitário", "Regras de alerta automático: IQAr, ruído diurno/noturno, perdas de água, cotas de córrego e leituras incompletas (esquema flexível)", "9"],
        ["`EquipamentoControllerTest`", "Unitário (camada web)", "Validação de entrada, códigos HTTP (201, 204, 400, 404, 409), formato Problem Details e conversão snake_case", "7"],
        ["`GaiaOpsApplicationIT`", "Integração", "Aplicação completa + **MongoDB 8.0 real (Testcontainers)**: carga inicial, painel ESG, ciclo CRUD com exclusão lógica, alerta automático, regra de negócio (422) e readiness", "6"],
    ], larguras=[4.3, 2.8, 8.2, 1.2])
    d.nota("A cobertura de código é medida pelo JaCoCo e publicada no artefato `relatorios-de-teste`. O resumo "
           "com aprovados, falhas e cobertura aparece na página de cada execução do pipeline (print 03).")


def secao_evidencias_pipeline(d):
    d.h1(SECOES[4])
    d.print_tela("01-pipeline-visao-geral", "Figura 3 - Execução completa do pipeline (5 jobs)",
                 "Actions > execução CI/CD: o grafo com Build, Testes, Imagem Docker, Deploy staging e Deploy producao, todos verdes.")
    d.print_tela("02-build", "Figura 4 - Job de build",
                 "Job '1. Build' com o passo 'Compilar e empacotar' expandido mostrando BUILD SUCCESS.")
    d.print_tela("03-testes", "Figura 5 - Resultado dos testes automatizados",
                 "Summary da execução: tabela 'Testes automatizados' com unitários, integração e cobertura.")
    d.print_tela("04-imagem-ghcr", "Figura 6 - Imagem publicada no GHCR",
                 "Packages > gaiaops-esg: tags do commit, latest, staging e production.")
    d.print_tela("05-deploy-staging", "Figura 7 - Deploy em staging e smoke test",
                 "Job '4. Deploy staging', passo 'Smoke test' expandido: 8/8 verificacoes OK.")
    d.print_tela("06-aprovacao-producao", "Figura 8 - Aprovação manual para produção",
                 "Execução aguardando revisão: aviso 'Waiting for review' / botão 'Review deployments'.")
    d.print_tela("07-deploy-producao", "Figura 9 - Deploy em produção e smoke test",
                 "Job '5. Deploy producao', passo 'Smoke test' expandido: 8/8 verificacoes OK.")


def secao_evidencias_ambientes(d):
    d.h1(SECOES[5])
    d.par("As capturas abaixo são geradas pelo próprio pipeline (`deploy/captura-tela.sh`, Chrome headless) "
          "logo após o smoke test de cada ambiente e publicadas nos artefatos `evidencias-staging` e "
          "`evidencias-production`. A faixa colorida no topo da página identifica o ambiente; os chips "
          "mostram a versão, o commit implantado e a saúde da aplicação.")
    d.print_tela("08-staging-pagina-inicial", "Figura 10 - Ambiente de staging funcionando",
                 "Artefato evidencias-staging/pagina-inicial.png.")
    d.print_tela("10-staging-swagger", "Figura 11 - Swagger UI disponível em staging",
                 "Artefato evidencias-staging/swagger-ui.png.")
    d.print_tela("09-producao-pagina-inicial", "Figura 12 - Ambiente de produção funcionando",
                 "Artefato evidencias-production/pagina-inicial.png.")
    d.print_tela("11-environments", "Figura 13 - Deployments por ambiente no GitHub",
                 "Página Deployments (ou Settings > Environments) com staging e production implantados.")


def secao_desafios(d):
    d.h1(SECOES[6])
    desafios = [
        ("Validar Docker e pipeline sem Docker na máquina de desenvolvimento",
         "A estação de desenvolvimento não tinha Docker instalado, então não dava para subir os containers "
         "localmente durante a construção.",
         "Validação em camadas: `docker compose config` (binário oficial, que não precisa do daemon) para "
         "conferir o Compose em cada ambiente; `hadolint` para o Dockerfile; `actionlint` e `shellcheck` para o "
         "workflow e os scripts; o mesmo `mongo/init/01-setup.js` executado no MongoDB local; e a API rodando "
         "contra ele com os perfis staging e prod, aprovada pelo mesmo smoke test do pipeline. Os testes de "
         "integração usam `@Testcontainers(disabledWithoutDocker = true)`: são ignorados sem Docker e sempre "
         "executados no pipeline."),
        ("Certificado TLS recusado pela JVM",
         "O Maven falhava com `PKIX path building failed` ao baixar dependências: a rede inspeciona o tráfego "
         "TLS com um certificado que o Java não reconhece (o `curl`, que usa o repositório do Windows, funcionava).",
         "Apontar a JVM para o repositório de certificados do Windows (`-Djavax.net.ssl.trustStoreType=Windows-ROOT` "
         "em `MAVEN_OPTS`), só na máquina local. O projeto não precisa de nenhuma alteração por causa disso."),
        ("Preservar o esquema flexível da Fase 5 na API",
         "Mapear os documentos para classes fixas faria um PUT apagar os campos específicos de cada tipo de "
         "equipamento, como os turnos de uma escola ou a vazão de uma ETE.",
         "A API trabalha com `org.bson.Document`. A atualização usa `$set` campo a campo, e os campos "
         "controlados não podem ser sobrescritos via atributos livres. O teste de integração comprova que um "
         "atributo não reenviado continua no documento."),
        ("Erros de validação saindo no formato padrão do Spring",
         "O tratador de Problem Details do Spring Boot tinha precedência sobre o da aplicação, e as mensagens de "
         "validação por campo se perdiam. Dois testes da camada web acusaram o problema.",
         "O tratador da aplicação passou a estender `ResponseEntityExceptionHandler`. O Spring Boot então "
         "desativa o dele, e todos os erros (inclusive JSON malformado) saem no mesmo formato, com mensagens em "
         "português."),
        ("Nome de campo inesperado no JSON",
         "Com a estratégia SNAKE_CASE do Jackson, o campo `licencasAVencer` virou `licencas_avencer`, e a página "
         "inicial não encontrava o valor. O teste manual da API revelou o problema.",
         "O campo foi renomeado para `licencasPorVencer` (`licencas_por_vencer`), e o teste de integração passou "
         "a verificar o nome."),
        ("Senhas do banco sem versionar segredos",
         "O deploy precisa de senhas, mas elas não podem ficar no repositório nem aparecer nos logs.",
         "Os arquivos `deploy/*.env` só têm configuração não sensível. As senhas vêm dos GitHub Secrets do "
         "environment ou são geradas no deploy; localmente, ficam em um arquivo ignorado pelo Git. Todas são "
         "mascaradas no log, e o Compose recusa subir sem elas (`${VAR:?}`)."),
        ("Docker Engine 29 nos runners e a API do Testcontainers",
         "Versões recentes do Docker Engine recusam clientes com API antiga, o que pode quebrar o Testcontainers.",
         "Spring Boot 3.5.16 (Testcontainers 1.21.4) e `src/test/resources/docker-java.properties` fixando a "
         "API mínima em 1.44."),
        ("Projeto editado no Windows, executado no Linux",
         "Scripts com fim de linha CRLF ou sem permissão de execução quebram no runner Linux.",
         "`.gitattributes` força LF em `.sh`, `mvnw`, `.yml` e `.env`, e os scripts são versionados com bit de "
         "execução (`git update-index --chmod=+x`)."),
        ("Garantir que produção receba exatamente o que foi testado",
         "Reconstruir a imagem a cada deploy poderia gerar artefatos diferentes entre staging e produção.",
         "Build once, deploy many: a imagem é construída uma vez, identificada pelo hash do commit, e promovida "
         "por tags. O smoke test confere o commit em `/actuator/info` em cada ambiente."),
    ]
    for i, (titulo, problema, solucao) in enumerate(desafios, 1):
        d.h2("7.%d %s" % (i, titulo))
        d.par("**Problema.** " + problema)
        d.par("**Solução.** " + solucao)


def secao_tecnologias(d):
    d.h1(SECOES[7])
    d.tabela(["Camada", "Tecnologias"], [
        ["Linguagem e framework", "Java 21, Spring Boot 3.5 (Web, Validation, Data MongoDB, Actuator), springdoc-openapi"],
        ["Banco de dados", "MongoDB 8.0 com validadores `$jsonSchema` e índices parciais, sparse, multikey, 2dsphere e de texto"],
        ["Build", "Maven 3.9 (Maven Wrapper)"],
        ["Testes", "JUnit 5, Spring MockMvc, Mockito, AssertJ, Testcontainers, JaCoCo"],
        ["Containers", "Docker (multi-stage), Docker Compose"],
        ["CI/CD", "GitHub Actions, GitHub Environments, GitHub Secrets, GitHub Container Registry"],
        ["Scripts e qualidade", "Bash, actionlint, shellcheck, hadolint"],
        ["Documentação", "Markdown, Mermaid, python-docx, LibreOffice"],
    ], larguras=[4.5, 12])


def secao_checklist(d):
    d.h1(SECOES[8])
    ambientes_ok = existe_print("08-staging-pagina-inicial") and existe_print("09-producao-pagina-inicial")
    prints_pipeline_ok = all(existe_print(n) for n in ("01-pipeline-visao-geral", "02-build", "03-testes",
                                                       "05-deploy-staging", "07-deploy-producao"))

    def marca(ok):
        return "☑" if ok else "☐"

    d.tabela(["Item", "OK"], [
        ["Projeto compactado em .ZIP com estrutura organizada", marca(True)],
        ["Dockerfile funcional", marca(os.path.exists(os.path.join(RAIZ, "Dockerfile")))],
        ["docker-compose.yml ou arquivos Kubernetes", marca(os.path.exists(os.path.join(RAIZ, "docker-compose.yml")))],
        ["Pipeline com etapas de build, teste e deploy",
         marca(os.path.exists(os.path.join(RAIZ, ".github", "workflows", "ci-cd.yml")))],
        ["README.md com instruções e prints", marca(prints_pipeline_ok and ambientes_ok)],
        ["Documentação técnica com evidências (PDF ou PPT)", marca(prints_pipeline_ok and ambientes_ok)],
        ["Deploy realizado nos ambientes staging e produção", marca(ambientes_ok)],
    ], larguras=[14, 2.5])


def main():
    d = Documento()
    capa(d)
    sumario(d)
    secao_projeto(d)
    secao_pipeline(d)
    secao_docker(d)
    secao_testes(d)
    secao_evidencias_pipeline(d)
    secao_evidencias_ambientes(d)
    secao_desafios(d)
    secao_tecnologias(d)
    secao_checklist(d)
    d.salvar()

    pendentes = [n for n in ("01-pipeline-visao-geral", "02-build", "03-testes", "04-imagem-ghcr",
                             "05-deploy-staging", "06-aprovacao-producao", "07-deploy-producao",
                             "08-staging-pagina-inicial", "09-producao-pagina-inicial",
                             "10-staging-swagger", "11-environments") if not existe_print(n)]
    print("Documento gerado: %s" % os.path.relpath(SAIDA_DOCX, RAIZ))
    if pendentes:
        print("Prints pendentes (%d) - veja docs/ROTEIRO-PRINTS.md:" % len(pendentes))
        for nome in pendentes:
            print("  - docs/prints/%s.png" % nome)
    else:
        print("Todos os prints presentes.")


if __name__ == "__main__":
    main()
